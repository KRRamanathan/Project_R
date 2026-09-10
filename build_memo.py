#!/usr/bin/env python3
"""Write MEMO.docx (2-page cream packet) from the locked headlines + figures."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import make_exec_charts

ROOT = Path(__file__).resolve().parent
NAVY = (0x1E, 0x27, 0x61)
TEAL = (0x1C, 0x72, 0x93)
BODY = (0x33, 0x33, 0x33)
MUTED = (0x8A, 0x8F, 0xA3)


def _shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _run(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.size = Pt(float(size))
    run.font.bold = bold
    run.font.name = name
    if color:
        run.font.color.rgb = RGBColor(*color)


def build_memo() -> Path:
    make_exec_charts.main()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    banner = doc.add_table(rows=1, cols=1)
    cell = banner.cell(0, 0)
    _shade(cell, "FFC400")
    p = cell.paragraphs[0]
    r = p.add_run("RAPIDO  ·  MEMO TO THE HEAD OF SUPPLY")
    _run(r, 10, True, NAVY)
    p2 = cell.add_paragraph()
    r2 = p2.add_run("Onboarding leak and overnight airport  ·  Ramanathan K R  ·  10 Sep 2026  ·  extract 30 Jun 2026, 23:59 IST")
    _run(r2, 9, False, NAVY)

    h = doc.add_paragraph()
    r = h.add_run("~180 more approved captains a month — two photo fixes, not a hiring push")
    _run(r, 16, True, NAVY, "Cambria")

    lead = doc.add_paragraph()
    r = lead.add_run(
        "The ask. Ship Insurance camera UX this week, then a 10-day RC grace after Legal. "
        "Do not hire the airport catchment. Stop scaling CAMP_WA_002. Medium confidence on 180; "
        "the campaign call is a hard no. These approvals are not hollow: 98.7% of mature approved "
        "captains already take a first trip (3,895 / 3,946)."
    )
    _run(r, 10.5, False, BODY)

    pics = doc.add_paragraph()
    pics.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for fname in ("pie_c1.png", "pie_airport_night.png"):
        pth = ROOT / "figures" / fname
        if pth.exists():
            pics.add_run().add_picture(str(pth), width=Inches(3.15))
            pics.add_run("  ")

    sub = doc.add_paragraph()
    r = sub.add_run("What to do this week")
    _run(r, 13, True, TEAL, "Cambria")

    tbl = doc.add_table(rows=4, cols=4)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Do this", "Impact", "Watch"]
    data = [
        [
            "1",
            "C1b Insurance camera this week (411), then C1a 10-day RC grace after Legal (1,637). Same camera on other docs is a free rider — not in the 180.",
            "~180/mo (~50 + ~134). C1a ops ~1 FTE ₹40–60k assumed.",
            "Approved by cohort",
        ],
        [
            "2",
            "Airport Return Assurance at ₹35.4 per unpaid night-suburban leg. Do not hire the catchment.",
            "Sample ~₹69–103k/mo. Rest-suburban, not city-core.",
            "Return-in-20m, cancels, falling 4-week run-rate",
        ],
        [
            "3",
            "Stop CAMP_WA_002 5×. Send/no-send RCT on RC-cleared app/paid captains only.",
            "0 pp click lift. Cost avoided + answer in 4–6 weeks.",
            "Aadhaar pass and approved",
        ],
    ]
    for i, htxt in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        _shade(cell, "1E2761")
        r = cell.paragraphs[0].add_run(htxt)
        _run(r, 9, True, (0xFF, 0xFF, 0xFF))
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri + 1].cells[ci]
            _shade(cell, "FFF8E1" if ri % 2 == 0 else "FFFDF6")
            r = cell.paragraphs[0].add_run(val)
            _run(r, 8.5, ci == 0, NAVY if ci == 0 else BODY)

    p = doc.add_paragraph()
    r = p.add_run(
        "Airport: ~40% unfulfilled vs ~3% elsewhere; 84% in 21:00–03:59. Mix does not shift overnight "
        "(χ² p=0.12). CAMP_WA_002 clicked vs not is −1.5pp (CI −3.6 to +0.6) → 0 pp. Do not fund 5×."
    )
    _run(r, 10.5, False, BODY)

    p = doc.add_paragraph()
    r = p.add_run("Why these two stages — leftover queue — ARA maths")
    _run(r, 13, True, TEAL, "Cambria")

    pics2 = doc.add_paragraph()
    pics2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for fname in ("pie_rc_capture.png", "pie_c1b.png"):
        pth = ROOT / "figures" / fname
        if pth.exists():
            pics2.add_run().add_picture(str(pth), width=Inches(3.15))
            pics2.add_run("  ")

    bar = ROOT / "figures" / "bar_volume_lost.png"
    if bar.exists():
        bp = doc.add_paragraph()
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bp.add_run().add_picture(str(bar), width=Inches(6.1))

    p = doc.add_paragraph()
    r = p.add_run(
        "RC loses 5,405; only 1,637 are capture-only (C1a). Insurance loses 3,289 but ~2,530 never uploaded "
        "after Fitness — C1b is the 411 photo fails. C1a: ~300/mo × 60% show-up × 74% pass ≈ 134. C1b: "
        "67→76% × 91% approved-if-cleared ≈ 46–52. If remaining docs bind, C1a shrinks toward ~36/month. "
        "ARA: gap ₹31.6 ÷ 0.8931 ≈ ₹35.4/leg. 30% of fare (₹105) overpays. Do not add field 128–256 "
        "(overlaps 726 C1a captains)."
    )
    _run(p.runs[0], 10.5, False, BODY)

    p = doc.add_paragraph()
    r = p.add_run("What I assumed, and what would change my answer. ")
    _run(r, 11, True, NAVY)
    r = p.add_run(
        "~15.6-day cutoff; doc_events not approvals.docs_cleared; field gap not banked; ARA vs rest-suburban "
        "not city-core; C1a FTE ₹40–60k assumed; trips file is sampled so ₹/month is not a city P&L."
    )
    _run(r, 10, False, BODY)

    out = ROOT / "MEMO.docx"
    doc.save(out)
    print(f"wrote {out}")
    return out


if __name__ == "__main__":
    build_memo()
